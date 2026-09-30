"""Download a CV and turn it into content Claude can read."""

import base64
import io
from pathlib import Path

import requests
from docx import Document


class UnreadableCV(Exception):
    pass


def download(url: str, dest: Path) -> Path:
    resp = requests.get(url, timeout=60)
    resp.raise_for_status()
    dest.write_bytes(resp.content)
    return dest


def to_content_block(path: Path) -> dict:
    """PDFs go to Claude as documents (it reads scanned PDFs too); Word files as extracted text."""
    ext = path.suffix.lower()
    data = path.read_bytes()
    if ext == ".pdf":
        return {
            "type": "document",
            "source": {"type": "base64", "media_type": "application/pdf", "data": base64.b64encode(data).decode()},
            "title": "CV",
        }
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(p for p in parts if p.strip())
        if not text.strip():
            raise UnreadableCV("Word CV contains no text")
        return {"type": "text", "text": f"<cv>\n{text}\n</cv>"}
    raise UnreadableCV(f"Unsupported CV format '{ext}' (only PDF and .docx)")
